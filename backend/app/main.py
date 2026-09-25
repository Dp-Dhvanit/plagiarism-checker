"""
FastAPI entrypoint for AI Text Detective.

Run (from backend/):
  uvicorn app.main:app --reload --port 8000
"""
from __future__ import annotations

import asyncio
import concurrent.futures
import io
import os
import time
from dataclasses import asdict
from typing import Any

import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import Response
from pydantic import BaseModel, Field

from app.scoring import BURSTINESS_WEIGHT, PERPLEXITY_WEIGHT, get_scorer
from app.humanizer import humanize_text
from app.code_analyzer import (
    analyze_code, detect_language,
    CodeSignals,
)
from app.summarizer import summarize_document, summarize_csv_bytes, _parse_csv_text
from app.gemini_summarizer import (
    build_prose_context,
    build_tabular_context,
    enhance_summary,
)
from app.text_quality import assess as assess_quality
from app.content_classifier import classify_block
from app.detection_pipeline import run as run_detection_pipeline, aggregate_code
from app.file_pipeline import build_units
from app.chunking import chunk_code
from app import db
from app.ai_text_detector import analyze_text_ai, GeminiTextDetection
from app.similarity import analyze_similarity, store_reference_chunks, SimilarityResult
from app.originality_rewriter import rewrite_to_reduce_overlap
from app.code_optimizer import optimize_code
from app.image_detector import analyze_image, validate_image, ImageValidationError
from app.pdf_report import generate_report

MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "15"))
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024

app = FastAPI(title="AI Text Detective", version="0.2.0")


def _env_list(name: str, default: str) -> list[str]:
    return [v.strip() for v in os.environ.get(name, default).split(",") if v.strip()]


# This API has no authentication and holds every analysed document, so it must
# not be readable by whatever web page happens to be open in the same browser.
# The bundled frontend reaches it through Vite's same-origin proxy, so it needs
# no CORS access at all; the allowed origins below only matter if the UI is
# served from somewhere else (set CORS_ORIGINS / ALLOWED_HOSTS to add them).
# The Host check is what stops DNS-rebinding, where a hostile page resolves its
# own domain to 127.0.0.1 and becomes "same-origin" with this server.
app.add_middleware(
    CORSMiddleware,
    allow_origins=_env_list("CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173"),
    allow_credentials=False,  # no cookies or sessions are used
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Content-Type"],
)
app.add_middleware(
    TrustedHostMiddleware,
    allowed_hosts=_env_list("ALLOWED_HOSTS", "localhost,127.0.0.1,testserver"),
)

# ── Source-code file extensions ──────────────────────────────────────────────
CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".java", ".cpp", ".c", ".cs",
    ".php", ".go", ".rb", ".kt", ".swift", ".sql",
}

DOCUMENT_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".docx", ".txt"}

# ══════════════════════════════════════════════════════════════════════════════
# Pydantic models
# ══════════════════════════════════════════════════════════════════════════════

class AnalyzeRequest(BaseModel):
    text: str = Field(..., min_length=1)
    ppl_weight: float | None = None
    burst_weight: float | None = None
    # Opt-in: also compare against Wikipedia / arXiv (app/web_sources.py).
    # Sends a few short excerpts of the text to those sites, so it is off
    # unless the user asks for it.
    check_web: bool = False


class SentenceBreakdown(BaseModel):
    sentence: str
    perplexity: float


class SignalsModel(BaseModel):
    perplexity_signal: float
    burstiness_signal: float
    marker_signal: float
    uniformity_signal: float


class CodeSignalsModel(BaseModel):
    comment_density: float
    generic_comment_ratio: float
    docstring_coverage: float
    naming_consistency: float
    error_handling_density: float
    avg_identifier_length: float
    func_length_uniformity: float
    # v2 signals
    cleanness_score: float
    edge_case_density: float
    line_uniformity: float


class CodeAnalyzeResponse(BaseModel):
    ai_code_probability: float
    human_code_probability: float
    confidence: str
    detected_language: str
    code_blocks_found: int
    signals: CodeSignalsModel
    explanation: str
    lines_analyzed: int
    # Present when the source was a multi-page/slide document analyzed as
    # several per-unit chunks (see /detect-code and detection_pipeline.py).
    chunks: list[dict] | None = None
    languages_detected: list[str] | None = None


class FlaggedSectionModel(BaseModel):
    text: str
    reason: str = ""


class GeminiTextModel(BaseModel):
    ai_probability: float
    confidence: str
    flagged_sections: list[FlaggedSectionModel] = []
    explanation: str


class SimilarityMatchModel(BaseModel):
    query_excerpt: str
    matched_text: str
    source_file: str
    # None for an external source, which has no history row.
    source_history_id: int | None = None
    score: float
    # Lexical corroboration for this specific pair — see SimilarityResult
    # docstring in app/similarity.py for why cosine alone isn't enough.
    word_overlap: float = 0.0
    ngram_overlap: float = 0.0
    verified: bool = False
    # Provenance: what the passage matched, and when/where it came from.
    # `source_kind` is "submission" (an earlier analysis on file, dated by
    # `source_created_at`) or "wikipedia" / "arxiv" (a page with a `source_url`).
    source_name: str = ""
    source_created_at: str | None = None
    source_kind: str = "submission"
    source_url: str | None = None


class ExternalSourceModel(BaseModel):
    kind: str
    title: str
    url: str


class ExternalCheckModel(BaseModel):
    """What the opt-in external-source check did. `sources` lists every page
    that was compared, matched or not."""

    status: str  # ok | no_candidates | unavailable | disabled | too_short
    note: str = ""
    sources: list[ExternalSourceModel] = []


class SimilarityModel(BaseModel):
    overall_similarity: float
    matches: list[SimilarityMatchModel] = []
    chunks_compared: int = 0
    corpus_size: int = 0
    note: str = ""
    # Strongest single passage match (0-100) — stable as the archive grows,
    # unlike the mean in `overall_similarity`. Raw embedding cosine: check
    # top_match_verified before treating this as a strong finding.
    top_match: float = 0.0
    top_match_verified: bool = False
    # Percentage of passages VERIFIED as copied (semantic + lexical
    # corroboration both present).
    matched_portion: float = 0.0
    # Percentage of passages that are semantically close but NOT verified —
    # topic overlap without shared phrasing.
    possible_portion: float = 0.0
    # Present only when the caller opted in to the external-source check.
    external: ExternalCheckModel | None = None


class DetectorResultModel(BaseModel):
    """One detection provider's independent opinion."""

    provider: str
    label: str = ""
    ai_probability: float | None = None
    verdict: str = "unavailable"
    confidence: str = "low"
    explanation: str = ""
    paid: bool = False
    local: bool = True
    error: str | None = None
    elapsed_ms: int = 0


class ConsensusModel(BaseModel):
    """How much the providers agreed. Deliberately not a blended score —
    averaging detectors that disagree would manufacture false confidence."""

    agreement: str = "none"       # unanimous | mixed | conflicted | none
    summary: str = ""
    providers_ran: int = 0
    min_score: float | None = None
    max_score: float | None = None
    spread: float | None = None


class AnalyzeResponse(BaseModel):
    # "analyzed" | "code" | "mixed" | "insufficient_text" | "unanalyzable" | "extraction_failure"
    status: str = "analyzed"
    message: str | None = None
    ai_likelihood_score: float
    perplexity: float
    burstiness: float
    verdict: str
    sentence_breakdown: list[SentenceBreakdown]
    signals: SignalsModel | None = None
    extracted_text: str | None = None
    code_result: CodeAnalyzeResponse | None = None
    document_class: str | None = None
    languages_detected: list[str] = []
    # Additive — Feature 1 (Gemini AI-probability + embedding-based
    # similarity). None when Gemini isn't configured/failed (gemini) or
    # simply not applicable (similarity is always populated when there is
    # analyzable prose text).
    gemini: GeminiTextModel | None = None
    similarity: SimilarityModel | None = None
    history_id: int | None = None
    # Additive — every registered detection provider that ran, plus how much
    # they agreed. `gemini` above is kept for backwards compatibility and is
    # also present here as one entry among several.
    detectors: list[DetectorResultModel] = []
    consensus: ConsensusModel | None = None
    # Additive — the verdict the UI should headline once every detector's
    # opinion is taken into account, and why. `verdict` above stays the local
    # detector's own verdict; this is "Uncertain" whenever the detectors that
    # ran disagree. None when no detector breakdown exists (e.g. code).
    final_verdict: str | None = None
    verdict_reason: str | None = None


class HumanizeRequest(BaseModel):
    text: str = Field(..., min_length=1)
    sentence_breakdown: list[SentenceBreakdown] = []
    ai_score: float = Field(default=50.0)


class HumanizeResponse(BaseModel):
    original: str
    humanized: str
    sentences_changed: int


class ImportantNumber(BaseModel):
    label: str
    value: str
    context: str = ""


class SummarizeResponse(BaseModel):
    document_type: str
    summary_points: list[str]
    word_count: int
    charts: list[dict] = []
    # Additive, optional — populated only when the Gemini semantic layer
    # ran successfully. Absent/empty means the response is exactly the
    # existing deterministic summary (Gemini not configured, or it failed).
    title: str | None = None
    overview: str | None = None
    important_numbers: list[ImportantNumber] = []
    ai_enhanced: bool = False


class ImageAnalyzeResponse(BaseModel):
    status: str = "analyzed"  # "analyzed" | "unavailable"
    message: str | None = None
    ai_probability: float = 0.0
    confidence: str = "low"
    classification: str = "uncertain"
    indicators: list[str] = []
    explanation: str = ""
    history_id: int | None = None
    # Which provider actually produced this result: "gemini" (primary) or
    # "openrouter" (fallback, only used when Gemini is unavailable/fails).
    provider: str | None = None


class HistoryItem(BaseModel):
    id: int
    file_name: str
    file_type: str
    analysis_type: str
    ai_probability: float | None = None
    similarity_score: float | None = None
    confidence: str | None = None
    status: str
    created_at: str
    report_path: str | None = None


class HistoryDetail(HistoryItem):
    result_json: dict[str, Any]


class DashboardStats(BaseModel):
    total_analyses: int
    text_analyses: int
    image_analyses: int
    ai_flagged: int
    average_similarity: float
    # Additive — week-over-week context for the overview cards. All optional
    # so older callers ignoring these fields are unaffected.
    average_ai_probability: float = 0.0
    scans_this_week: int = 0
    scans_prev_week: int = 0
    avg_ai_this_week: float | None = None
    avg_ai_prev_week: float | None = None
    avg_similarity_this_week: float | None = None
    avg_similarity_prev_week: float | None = None


# ══════════════════════════════════════════════════════════════════════════════
# Startup
# ══════════════════════════════════════════════════════════════════════════════

@app.on_event("startup")
def warmup_model() -> None:
    get_scorer().load()
    db.init_db()


@app.get("/health")
def health() -> dict:
    return {"status": "ok"}


# ══════════════════════════════════════════════════════════════════════════════
# Text extraction helpers
# ══════════════════════════════════════════════════════════════════════════════

def _extract_pdf_text(data: bytes) -> str:
    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                parts.append(t.strip())
    return "\n\n".join(parts)


def _extract_pptx_text(data: bytes, join_bullets: bool = True) -> str:
    prs = Presentation(io.BytesIO(data))
    slides_text: list[str] = []
    for slide in prs.slides:
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs).strip()
                    if line:
                        slide_lines.append(line)
        if slide_lines:
            if join_bullets:
                slides_text.append(". ".join(slide_lines) + ".")
            else:
                slides_text.append("\n".join(slide_lines))
    return " ".join(slides_text) if join_bullets else "\n\n".join(slides_text)


def _extract_docx_text(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_by_filename(data: bytes, filename: str, join_bullets: bool = True) -> str:
    fn = filename.lower()
    if fn.endswith(".pdf"):
        return _extract_pdf_text(data)
    if fn.endswith((".pptx", ".ppt")):
        return _extract_pptx_text(data, join_bullets=join_bullets)
    if fn.endswith(".docx"):
        return _extract_docx_text(data)
    # TXT or code file
    return data.decode("utf-8", errors="ignore")


# ══════════════════════════════════════════════════════════════════════════════
# Existing endpoints (unchanged)
# ══════════════════════════════════════════════════════════════════════════════

def _build_response(
    result,
    extracted_text: str | None = None,
    status: str = "analyzed",
    message: str | None = None,
    code_result: CodeAnalyzeResponse | None = None,
    document_class: str | None = None,
    languages_detected: list[str] | None = None,
    gemini: GeminiTextDetection | None = None,
    similarity: SimilarityResult | None = None,
    history_id: int | None = None,
    detectors: list | None = None,
) -> AnalyzeResponse:
    signals = None
    if result.signals:
        signals = SignalsModel(
            perplexity_signal=result.signals.perplexity_signal,
            burstiness_signal=result.signals.burstiness_signal,
            marker_signal=result.signals.marker_signal,
            uniformity_signal=result.signals.uniformity_signal,
        )
    head = _headline_for(detectors, result.verdict)
    return AnalyzeResponse(
        status=status,
        message=message,
        ai_likelihood_score=result.ai_likelihood_score,
        perplexity=result.perplexity,
        burstiness=result.burstiness,
        verdict=result.verdict,
        sentence_breakdown=[
            SentenceBreakdown(sentence=s.sentence, perplexity=s.perplexity)
            for s in result.sentence_breakdown
        ],
        signals=signals,
        extracted_text=extracted_text,
        code_result=code_result,
        document_class=document_class,
        languages_detected=languages_detected or [],
        gemini=_gemini_to_model(gemini),
        similarity=_similarity_to_model(similarity),
        history_id=history_id,
        detectors=_detectors_to_models(detectors),
        consensus=_consensus_to_model(detectors),
        final_verdict=head["verdict"] if head else None,
        verdict_reason=head["reason"] if head else None,
    )


# ══════════════════════════════════════════════════════════════════════════════
# NEW: Feature 1 helpers — Gemini AI-probability + embedding similarity,
# and logging completed text analyses into history (Feature 3).
# ══════════════════════════════════════════════════════════════════════════════

def _detectors_to_models(results: list | None) -> list[DetectorResultModel]:
    """Serialise every provider's opinion, including ones that couldn't run —
    a skipped provider is information, not something to hide."""
    if not results:
        return []
    from app.detectors.registry import _REGISTRY

    meta = {d.name: d for d in _REGISTRY}
    out = []
    for r in results:
        d = meta.get(r.provider)
        out.append(
            DetectorResultModel(
                provider=r.provider,
                label=d.label if d else r.provider,
                ai_probability=r.ai_probability,
                verdict=r.verdict,
                confidence=r.confidence,
                explanation=r.explanation,
                paid=bool(d.paid) if d else False,
                local=bool(d.local) if d else True,
                error=r.error,
                elapsed_ms=r.elapsed_ms,
            )
        )
    return out


def _consensus_to_model(results: list | None) -> ConsensusModel | None:
    if not results:
        return None
    from app.detectors.registry import consensus

    c = consensus(results)
    return ConsensusModel(
        agreement=c.get("agreement", "none"),
        summary=c.get("summary", ""),
        providers_ran=c.get("providers_ran", 0),
        min_score=c.get("min_score"),
        max_score=c.get("max_score"),
        spread=c.get("spread"),
    )


def _headline_for(results: list | None, local_verdict: str) -> dict | None:
    """The verdict to headline given every detector's opinion (see
    detectors.registry.headline). None when there is no breakdown to judge."""
    if not results:
        return None
    from app.detectors.registry import headline

    return headline(results, local_label=local_verdict)


def _gemini_to_model(g: GeminiTextDetection | None) -> GeminiTextModel | None:
    if g is None:
        return None
    return GeminiTextModel(
        ai_probability=g.ai_probability,
        confidence=g.confidence,
        flagged_sections=[FlaggedSectionModel(text=s.text, reason=s.reason) for s in g.flagged_sections],
        explanation=g.explanation,
    )


def _similarity_to_model(s: SimilarityResult | None) -> SimilarityModel | None:
    if s is None:
        return None
    return SimilarityModel(
        overall_similarity=s.overall_similarity,
        matches=[SimilarityMatchModel(**asdict(m)) for m in s.matches],
        chunks_compared=s.chunks_compared,
        corpus_size=s.corpus_size,
        note=s.note,
        top_match=s.top_match,
        top_match_verified=s.top_match_verified,
        matched_portion=s.matched_portion,
        possible_portion=s.possible_portion,
        external=ExternalCheckModel(**s.external) if s.external else None,
    )


def _derive_confidence(score: float) -> str:
    """Fallback confidence label when Gemini isn't configured — how
    decisive the local heuristic score is (far from the 50/50 midpoint),
    not a claim about accuracy."""
    distance = abs(score - 50.0)
    if distance >= 30:
        return "high"
    if distance >= 15:
        return "medium"
    return "low"


def _assemble_detectors(
    prose_result, gemini_result, gemini_ms: int, hosted: list
) -> list:
    """Every detector's opinion, in registry order.

    Two of them are not re-run: the local heuristic is rebuilt from the
    `prose_result` we already have (re-running it would recompute perplexity
    over the whole document), and Gemini is rebuilt from the single call made
    for the legacy panel (a second call for the same text would spend another
    of the free tier's 20 daily requests). `hosted` is everything else.
    """
    from app.detectors.base import DetectorResult
    from app.detectors.registry import get_detector, in_registry_order

    local = DetectorResult(
        provider="heuristic",
        ai_probability=prose_result.ai_likelihood_score,
        verdict=DetectorResult.verdict_for(prose_result.ai_likelihood_score),
        confidence="low",
        explanation=(
            f"Perplexity {prose_result.perplexity}, burstiness {prose_result.burstiness}. "
            "Derived from writing-pattern statistics only."
        ),
    )
    opinions = [local, *hosted]
    gemini_detector = get_detector("gemini")
    if gemini_detector is not None:  # someone may have de-registered it; don't 500 over that
        gemini = gemini_detector.from_detection(gemini_result)
        gemini.elapsed_ms = gemini_ms
        opinions.append(gemini)
    return in_registry_order(opinions)


def _compare_corpus(
    text: str, exclude_file_name: str | None, check_web: bool
) -> SimilarityResult:
    """Similarity against earlier submissions, plus — only if asked — the
    Wikipedia/arXiv pages found for this text (fetched first, since their
    passages have to be in hand before the comparison runs)."""
    external = None
    if check_web:
        from app.web_sources import check_external

        external = check_external(text)
    result = analyze_similarity(
        text, exclude_file_name, extra_corpus=external.corpus if external else None
    )
    if external is not None:
        result.external = external.to_dict()
    return result


def _run_text_ai_features(
    text: str,
    exclude_file_name: str | None = None,
    prose_result=None,
    check_web: bool = False,
) -> tuple[GeminiTextDetection | None, SimilarityResult, list]:
    """The checks that follow local scoring — Gemini, corpus similarity and
    the other hosted detectors — run concurrently. They are independent, and
    the slowest (a hosted model call) used to be waited on before the next
    one even started."""
    from app.detectors.registry import run_all

    def timed_gemini():
        t0 = time.perf_counter()
        return analyze_text_ai(text), int((time.perf_counter() - t0) * 1000)

    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        gemini_f = pool.submit(timed_gemini)
        similarity_f = pool.submit(_compare_corpus, text, exclude_file_name, check_web)
        hosted_f = (
            pool.submit(run_all, text, exclude={"heuristic", "gemini"})
            if prose_result is not None
            else None
        )
        gemini_result, gemini_ms = gemini_f.result()
        similarity_result = similarity_f.result()
        hosted = hosted_f.result() if hosted_f is not None else []

    detectors = (
        _assemble_detectors(prose_result, gemini_result, gemini_ms, hosted)
        if prose_result is not None
        else []
    )
    return gemini_result, similarity_result, detectors


async def _run_text_ai_features_async(
    text: str,
    exclude_file_name: str | None = None,
    prose_result=None,
    check_web: bool = False,
) -> tuple[GeminiTextDetection | None, SimilarityResult, list]:
    return await asyncio.to_thread(
        _run_text_ai_features, text, exclude_file_name, prose_result, check_web
    )


def _pasted_text_tag(text: str) -> str:
    """Deterministic per-content identifier for pasted text.

    Used to tag stored corpus chunks, and by /reduce-overlap to exclude the
    copy of this exact text that /analyze stored moments earlier — never shown
    as the History list's display name (that stays the readable "Pasted
    text"). Every pasted submission shares that literal display string, so
    without a per-content tag they can't be told apart.

    /analyze deliberately does NOT exclude by this tag: it compares before it
    stores, so there is no self-match to avoid, and excluding would make an
    identical paste from someone else invisible — the one thing a plagiarism
    check must catch.
    """
    import hashlib

    digest = hashlib.sha256(text.strip().encode("utf-8")).hexdigest()[:12]
    return f"Pasted text ({digest})"


def _log_text_history(
    file_name: str,
    file_type: str,
    prose_result,
    gemini_result: GeminiTextDetection | None,
    similarity_result: SimilarityResult,
    full_text: str,
    corpus_tag: str | None = None,
    detectors: list | None = None,
) -> int:
    ai_probability = gemini_result.ai_probability if gemini_result else prose_result.ai_likelihood_score
    confidence = gemini_result.confidence if gemini_result else _derive_confidence(prose_result.ai_likelihood_score)

    detector_models = _detectors_to_models(detectors)
    consensus_model = _consensus_to_model(detectors)
    head = _headline_for(detectors, prose_result.verdict)

    result_json: dict[str, Any] = {
        "heuristic": {
            "ai_likelihood_score": prose_result.ai_likelihood_score,
            "verdict": prose_result.verdict,
            "perplexity": prose_result.perplexity,
            "burstiness": prose_result.burstiness,
        },
        "gemini_text": gemini_result.model_dump() if gemini_result else None,
        "similarity": {
            "overall_similarity": similarity_result.overall_similarity,
            "matches": [asdict(m) for m in similarity_result.matches],
            "chunks_compared": similarity_result.chunks_compared,
            "corpus_size": similarity_result.corpus_size,
            "note": similarity_result.note,
            "top_match": similarity_result.top_match,
            "top_match_verified": similarity_result.top_match_verified,
            "matched_portion": similarity_result.matched_portion,
            "possible_portion": similarity_result.possible_portion,
            "external": similarity_result.external,
        },
        "image_result": None,
        "extracted_text_preview": full_text[:3000],
        # Additive — every registered provider's independent opinion plus
        # how much they agreed, so it survives past the single live
        # response and can be shown again later (e.g. the Overview page's
        # "AI Detection Breakdown"). Not exposed anywhere before this.
        "detectors": [d.model_dump() for d in detector_models],
        "consensus": consensus_model.model_dump() if consensus_model else None,
        "final_verdict": head["verdict"] if head else None,
        "verdict_reason": head["reason"] if head else None,
    }
    history_id = db.insert_history(
        file_name=file_name,
        file_type=file_type,
        analysis_type="text",
        ai_probability=round(ai_probability, 1),
        # Verified-overlap portion, not raw top-passage cosine: a document
        # that merely shares a topic with something on file can score high
        # on semantics alone (see app/similarity.py), and that must not
        # inflate the History list or the dashboard's average similarity.
        similarity_score=similarity_result.matched_portion,
        confidence=confidence,
        status="analyzed",
        result_json=result_json,
    )
    store_reference_chunks(history_id, corpus_tag or file_name, full_text)
    return history_id


def _neutral_response(
    status: str,
    message: str | None,
    extracted_text: str | None = None,
    code_result: CodeAnalyzeResponse | None = None,
    document_class: str | None = None,
    languages_detected: list[str] | None = None,
) -> AnalyzeResponse:
    """Response for statuses that never ran (or only partially ran) the
    prose scorer — carries neutral placeholders instead of a fabricated
    score, per the "don't manufacture an AI score" requirement."""
    return AnalyzeResponse(
        status=status,
        message=message,
        ai_likelihood_score=0.0,
        perplexity=0.0,
        burstiness=0.0,
        verdict="Uncertain",
        sentence_breakdown=[],
        signals=None,
        extracted_text=extracted_text,
        code_result=code_result,
        document_class=document_class,
        languages_detected=languages_detected or [],
    )


@app.post("/analyze", response_model=AnalyzeResponse)
def analyze(body: AnalyzeRequest) -> AnalyzeResponse:
    text = body.text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="Please provide some text.")

    content_class = classify_block(text)

    if content_class in ("code", "mixed"):
        code_result = analyze_code(text, filename="")
        code_response = _build_code_response(code_result)

        if content_class == "code":
            return _neutral_response(status="code", message=None, code_result=code_response)

        # Mixed content: still attempt prose scoring on the whole text as a
        # best-effort second signal alongside the code result.
        quality = assess_quality(text)
        if quality.status == "ok":
            prose_result = get_scorer().analyze(text)
            tag = _pasted_text_tag(text)
            gemini_result, similarity_result, detectors = _run_text_ai_features(
                text, prose_result=prose_result, check_web=body.check_web
            )
            history_id = _log_text_history(
                "Pasted text", "text", prose_result, gemini_result, similarity_result, text,
                corpus_tag=tag, detectors=detectors,
            )
            return _build_response(
                prose_result, status="mixed", code_result=code_response,
                gemini=gemini_result, similarity=similarity_result, history_id=history_id,
                detectors=detectors,
            )
        return _neutral_response(status="mixed", message=quality.reason, code_result=code_response)

    quality = assess_quality(text)
    if quality.status != "ok":
        return _neutral_response(status=quality.status, message=quality.reason)

    scorer = get_scorer()
    result = scorer.analyze(text)
    tag = _pasted_text_tag(text)
    gemini_result, similarity_result, detectors = _run_text_ai_features(
        text, prose_result=result, check_web=body.check_web
    )
    history_id = _log_text_history(
        "Pasted text", "text", result, gemini_result, similarity_result, text,
        corpus_tag=tag, detectors=detectors,
    )
    return _build_response(
        result, gemini=gemini_result, similarity=similarity_result,
        history_id=history_id, detectors=detectors,
    )


@app.post("/upload", response_model=AnalyzeResponse)
async def upload_file(
    file: UploadFile = File(...),
    check_web: bool = Form(False),  # opt-in Wikipedia/arXiv comparison; see AnalyzeRequest
) -> AnalyzeResponse:
    filename = (file.filename or "").lower()
    if not any(filename.endswith(ext) for ext in DOCUMENT_EXTENSIONS):
        raise HTTPException(status_code=400, detail="Only PDF, PPTX, PPT, DOCX, or TXT files are supported.")
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File is too large ({len(data) / (1024 * 1024):.1f} MB). Maximum allowed size is {MAX_UPLOAD_MB} MB.",
        )
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    # PDF/DOCX/PPTX parsing is CPU-bound and synchronous; run it off the
    # event loop so a large upload doesn't stall every other in-flight
    # request while it parses.
    pipeline = await asyncio.to_thread(run_detection_pipeline, filename, data)
    code_response = _build_code_response(pipeline.code_result) if pipeline.code_result else None
    preview = pipeline.extracted_preview or None

    if pipeline.text_result:
        # Exclude this same filename so re-uploading a document doesn't
        # report it as plagiarising its own earlier submission.
        gemini_result, similarity_result, detectors = await _run_text_ai_features_async(
            pipeline.full_prose_text,
            file.filename or filename,
            prose_result=pipeline.text_result,
            check_web=check_web,
        )
        history_id = _log_text_history(
            file.filename or filename, filename.rsplit(".", 1)[-1],
            pipeline.text_result, gemini_result, similarity_result, pipeline.full_prose_text,
            detectors=detectors,
        )
        return _build_response(
            pipeline.text_result,
            extracted_text=preview,
            status=pipeline.status,
            message=pipeline.message,
            code_result=code_response,
            document_class=pipeline.document_class,
            languages_detected=pipeline.languages_detected,
            gemini=gemini_result,
            similarity=similarity_result,
            history_id=history_id,
            detectors=detectors,
        )

    return _neutral_response(
        status=pipeline.status,
        message=pipeline.message,
        extracted_text=preview,
        code_result=code_response,
        document_class=pipeline.document_class,
        languages_detected=pipeline.languages_detected,
    )


@app.post("/humanize", response_model=HumanizeResponse)
def humanize(body: HumanizeRequest) -> HumanizeResponse:
    text = body.text.strip()
    breakdown = [{"sentence": s.sentence, "perplexity": s.perplexity} for s in body.sentence_breakdown]
    humanized, changed = humanize_text(text, breakdown, ai_score=body.ai_score)
    return HumanizeResponse(original=text, humanized=humanized, sentences_changed=changed)


# ══════════════════════════════════════════════════════════════════════════════
# NEW: Similarity-aware rewrite (distinct from /humanize above, which targets
# the local AI-detector score, not corpus overlap)
# ══════════════════════════════════════════════════════════════════════════════

class ReduceOverlapRequest(BaseModel):
    text: str = Field(..., min_length=1)
    # Excludes a prior submission with this filename from the comparison,
    # matching /analyze's self-match prevention for re-submitted documents.
    file_name: str | None = None
    # Gemini is a candidate-generation source here too (see
    # app/originality_rewriter.py), but never called unless explicitly
    # requested — its free quota is scarce and this endpoint already tries
    # the free local rewriter, Groq, and OpenRouter first.
    include_gemini: bool = False


class ReduceOverlapResponse(BaseModel):
    original: str
    rewritten: str
    changed: bool
    improved: bool
    attempts_tried: int
    sentences_rewritten: int
    before: SimilarityModel
    after: SimilarityModel
    note: str
    # Which candidate source actually won: "local" | "groq" | "openrouter" |
    # "gemini" | "none". Transparency only — the similarity numbers above
    # are what actually decided the result.
    source: str = "none"
    sources_tried: list[str] = []


@app.post("/reduce-overlap", response_model=ReduceOverlapResponse)
def reduce_overlap(body: ReduceOverlapRequest) -> ReduceOverlapResponse:
    """Rewrite only the passages VERIFIED as overlapping with stored
    documents, then confirm the rewrite actually reduced that overlap
    before returning it. See app/originality_rewriter.py for the full
    pipeline and the measured evidence behind it."""
    text = body.text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="Please provide some text.")
    # If the caller didn't identify a real source file, fall back to the
    # same content-hash tag /analyze uses for pasted text — otherwise a
    # "reduce overlap" call on text just analyzed a moment ago would
    # self-match against the copy /analyze had just stored.
    exclude = body.file_name or _pasted_text_tag(text)
    result = rewrite_to_reduce_overlap(text, exclude_file_name=exclude, include_gemini=body.include_gemini)
    return ReduceOverlapResponse(
        original=result.original_text,
        rewritten=result.rewritten_text,
        changed=result.changed,
        improved=result.improved,
        attempts_tried=result.attempts_tried,
        sentences_rewritten=result.sentences_rewritten,
        before=_similarity_to_model(result.before),
        after=_similarity_to_model(result.after),
        note=result.note,
        source=result.source,
        sources_tried=result.sources_tried,
    )


# ══════════════════════════════════════════════════════════════════════════════
# NEW: AI Code Detection
# ══════════════════════════════════════════════════════════════════════════════

class CodeAnalyzeRequest(BaseModel):
    code: str = Field(..., min_length=10, description="Raw source code to analyze")
    filename: str = Field(default="", description="Optional filename for language hint")


def _build_code_response(result) -> CodeAnalyzeResponse:
    return CodeAnalyzeResponse(
        ai_code_probability=result.ai_code_probability,
        human_code_probability=result.human_code_probability,
        confidence=result.confidence,
        detected_language=result.detected_language,
        code_blocks_found=result.code_blocks_found,
        signals=CodeSignalsModel(
            comment_density=result.signals.comment_density,
            generic_comment_ratio=result.signals.generic_comment_ratio,
            docstring_coverage=result.signals.docstring_coverage,
            naming_consistency=result.signals.naming_consistency,
            error_handling_density=result.signals.error_handling_density,
            avg_identifier_length=result.signals.avg_identifier_length,
            func_length_uniformity=result.signals.func_length_uniformity,
            cleanness_score=result.signals.cleanness_score,
            edge_case_density=result.signals.edge_case_density,
            line_uniformity=result.signals.line_uniformity,
        ),
        explanation=result.explanation,
        lines_analyzed=result.lines_analyzed,
    )


@app.post("/detect-code", response_model=CodeAnalyzeResponse)
async def detect_code_file(file: UploadFile = File(...)) -> CodeAnalyzeResponse:
    filename = file.filename or ""
    fn_lower = filename.lower()
    data = await file.read()

    ext = "." + fn_lower.rsplit(".", 1)[-1] if "." in fn_lower else ""

    try:
        if ext in CODE_EXTENSIONS:
            # Direct source file — read as text
            code_text = data.decode("utf-8", errors="ignore")
            result = analyze_code(code_text, filename=filename)
            return _build_code_response(result)

        elif any(fn_lower.endswith(e) for e in DOCUMENT_EXTENSIONS):
            # Document — extract page/slide-aware units so multi-page
            # documents with code on different pages (possibly in
            # different languages) are tracked and aggregated per-unit,
            # instead of flattened into one blob. Uses classify_block()
            # (ratio-based over the whole unit) rather than
            # extract_code_blocks()'s stricter "4 consecutive matching
            # lines" rule, which can silently drop a short code page/slide
            # whose lines are broken up by bare brace lines.
            units = build_units(filename, data)
            non_empty = [u for u in units if u.text.strip()]
            if not non_empty:
                raise HTTPException(status_code=422, detail="No readable text found in the document.")

            code_units = [u for u in non_empty if classify_block(u.text) in ("code", "mixed")]
            if not code_units:
                raise HTTPException(
                    status_code=422,
                    detail="No programming code detected in this document. "
                           "Make sure the document contains selectable (non-image) code.",
                )

            code_chunks = chunk_code(code_units)
            agg_result, languages = aggregate_code(code_chunks)
            response = _build_code_response(agg_result)
            response.chunks = [
                {"units": c.unit_indices, "lines": c.line_count}
                for c in code_chunks
            ]
            response.languages_detected = languages if len(languages) > 1 else None
            return response

        else:
            # Try to read as plain text / code
            code_text = data.decode("utf-8", errors="ignore")
            result = analyze_code(code_text, filename=filename)
            return _build_code_response(result)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Analysis failed: {exc}")


@app.post("/detect-code-text", response_model=CodeAnalyzeResponse)
def detect_code_text(body: CodeAnalyzeRequest) -> CodeAnalyzeResponse:
    """Analyze raw code pasted as text."""
    code = body.code.strip()
    if len(code) < 10:
        raise HTTPException(status_code=400, detail="Please provide at least 10 characters of code.")
    result = analyze_code(code, filename=body.filename)
    return _build_code_response(result)


# ══════════════════════════════════════════════════════════════════════════════
# NEW: Code Optimizer — an addition to Code Analysis, not a replacement.
# Uses the SAME code the user already submitted for analysis (see
# app/code_optimizer.py); never re-extracts or re-uploads anything.
# ══════════════════════════════════════════════════════════════════════════════

class OptimizeCodeRequest(BaseModel):
    code: str = Field(..., min_length=1)
    language: str | None = None
    # Gemini is a candidate source here too, but — same as /reduce-overlap —
    # never called unless explicitly requested, to protect its free quota.
    include_gemini: bool = False


class LocalCodeMetricsModel(BaseModel):
    lines: int
    functions: int
    classes: int
    imports: int


class ComplexitySideModel(BaseModel):
    time: str = "O(?)"
    space: str = "O(?)"


class ComplexityEstimateModel(BaseModel):
    # The MODEL's own claim, never locally verified — labeled as an
    # estimate in the UI, not a measured fact.
    before: ComplexitySideModel = ComplexitySideModel()
    after: ComplexitySideModel = ComplexitySideModel()
    reasoning: str = ""
    # A genuinely local cross-check (real AST parse), Python only. None for
    # every other language rather than a guessed/unreliable number.
    loop_nesting_before: int | None = None
    loop_nesting_after: int | None = None


class OptimizeCodeResponse(BaseModel):
    available: bool
    provider: str | None = None
    optimized_code: str
    changed: bool
    changes: list[str] = []
    summary: str = ""
    original: LocalCodeMetricsModel
    optimized: LocalCodeMetricsModel
    validation_passed: bool
    validation_label: str
    validation_notes: list[str] = []
    complexity: ComplexityEstimateModel
    error: str | None = None


@app.post("/optimize-code", response_model=OptimizeCodeResponse)
def optimize_code_endpoint(body: OptimizeCodeRequest) -> OptimizeCodeResponse:
    code = body.code.strip()
    if len(code) < 10:
        raise HTTPException(status_code=400, detail="Please provide at least 10 characters of code.")
    result = optimize_code(code, language=body.language or "unknown", include_gemini=body.include_gemini)
    return OptimizeCodeResponse(
        available=result.available,
        provider=result.provider,
        optimized_code=result.optimized_code,
        changed=result.changed,
        changes=result.changes,
        summary=result.summary,
        original=LocalCodeMetricsModel(**vars(result.original)),
        optimized=LocalCodeMetricsModel(**vars(result.optimized)),
        validation_passed=result.validation_passed,
        validation_label=result.validation_label,
        validation_notes=result.validation_notes,
        complexity=ComplexityEstimateModel(
            before=ComplexitySideModel(**vars(result.complexity.before)),
            after=ComplexitySideModel(**vars(result.complexity.after)),
            reasoning=result.complexity.reasoning,
            loop_nesting_before=result.complexity.loop_nesting_before,
            loop_nesting_after=result.complexity.loop_nesting_after,
        ),
        error=result.error,
    )


# ══════════════════════════════════════════════════════════════════════════════
# NEW: Smart Document Summary
# ══════════════════════════════════════════════════════════════════════════════

# Maps Gemini's chart-type recommendation onto the existing chart_selector
# type vocabulary. Gemini only ever picks AMONG charts the deterministic
# pipeline already built from real data — it never supplies chart values.
_GEMINI_CHART_TYPE_MAP = {"bar": "bar", "horizontal_bar": "hbar", "pie": "pie", "donut": "pie", "line": "line"}


def _resolve_charts(gemini_result, deterministic_charts: list[dict]) -> list[dict]:
    if gemini_result is None:
        return deterministic_charts
    viz = gemini_result.visualization
    if not viz.recommended or viz.chart_type == "none":
        return []
    wanted = _GEMINI_CHART_TYPE_MAP.get(viz.chart_type)
    if wanted:
        matched = [c for c in deterministic_charts if c.get("type") == wanted]
        if matched:
            return matched[:1]
    # "scatter" (not implemented) or a type the deterministic pipeline
    # didn't produce -> keep the deterministic default rather than
    # showing nothing or fabricating something.
    return deterministic_charts


async def _apply_gemini(base: SummarizeResponse, context: dict, kind: str, charts: list[dict]) -> SummarizeResponse:
    """Best-effort semantic enhancement. Any failure anywhere in this
    function must fall back to `base` (the existing deterministic
    response) unchanged — never raise, never crash the request."""
    try:
        gemini_result = await enhance_summary(context, kind)
    except Exception:
        gemini_result = None

    if gemini_result is None:
        return base

    base.title = gemini_result.title
    base.overview = gemini_result.overview
    base.summary_points = gemini_result.key_takeaways or base.summary_points
    base.important_numbers = [
        ImportantNumber(label=n.label, value=n.value, context=n.context)
        for n in gemini_result.important_numbers
    ]
    base.charts = _resolve_charts(gemini_result, charts)
    base.ai_enhanced = True
    return base


@app.post("/summarize", response_model=SummarizeResponse)
async def summarize(file: UploadFile = File(...)) -> SummarizeResponse:
    filename = (file.filename or "").lower()
    allowed = DOCUMENT_EXTENSIONS | {".txt", ".csv", ".xlsx", ".xls"}
    if not any(filename.endswith(ext) for ext in allowed):
        raise HTTPException(
            status_code=400,
            detail="Supported formats: PDF, PPTX, DOCX, TXT, CSV.",
        )
    data = await file.read()

    # CSV / Excel — treat as tabular data
    if filename.endswith((".csv", ".xlsx", ".xls")):
        doc_type, points, charts = summarize_csv_bytes(data)
        response = SummarizeResponse(
            document_type="Dataset / Spreadsheet",
            summary_points=points,
            word_count=len(data.decode("utf-8", errors="ignore").split()),
            charts=charts,
        )
        try:
            headers, rows = _parse_csv_text(data.decode("utf-8", errors="ignore"))
            context = build_tabular_context(filename, headers, rows, points)
        except Exception:
            return response
        return await _apply_gemini(response, context, "tabular", charts)

    try:
        text = await asyncio.to_thread(_extract_by_filename, data, filename, join_bullets=False)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not extract text: {exc}")

    text = text.strip()
    if len(text) < 50:
        raise HTTPException(
            status_code=422,
            detail="Document contains too little readable text to summarize.",
        )

    doc_type, points, charts = summarize_document(text, n_points=8)
    word_count = len(text.split())

    response = SummarizeResponse(
        document_type=doc_type.replace("_", " ").title(),
        summary_points=points,
        word_count=word_count,
        charts=charts,
    )
    try:
        context = build_prose_context(filename, doc_type, points, word_count)
    except Exception:
        return response
    return await _apply_gemini(response, context, "prose", charts)


# ══════════════════════════════════════════════════════════════════════════════
# NEW: AI Image Detection (Feature 2)
# ══════════════════════════════════════════════════════════════════════════════

ALLOWED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


@app.post("/analyze/image", response_model=ImageAnalyzeResponse)
async def analyze_image_endpoint(file: UploadFile = File(...)) -> ImageAnalyzeResponse:
    filename = file.filename or "image"
    fn_lower = filename.lower()
    if not any(fn_lower.endswith(ext) for ext in ALLOWED_IMAGE_EXTENSIONS):
        raise HTTPException(status_code=400, detail="Only JPG, JPEG, PNG, or WEBP images are supported.")

    data = await file.read()
    try:
        mime = validate_image(data)
    except ImageValidationError as exc:
        raise HTTPException(status_code=422, detail=str(exc))

    result, reason = await asyncio.to_thread(analyze_image, data, mime)
    if result is None:
        # The full exception is already in the server log (see
        # gemini_client.log_failure); `reason` is the short, key-safe
        # version so the UI can say what actually went wrong.
        return ImageAnalyzeResponse(
            status="unavailable",
            message=f"AI image detection is unavailable — {reason}"
            if reason
            else "AI image detection is currently unavailable on the server. Please try again later.",
        )

    result_json: dict[str, Any] = {
        "heuristic": None,
        "gemini_text": None,
        "similarity": None,
        "image_result": result.model_dump(),
        "extracted_text_preview": None,
    }
    history_id = db.insert_history(
        file_name=filename,
        file_type=fn_lower.rsplit(".", 1)[-1],
        analysis_type="image",
        ai_probability=round(result.ai_probability, 1),
        similarity_score=None,
        confidence=result.confidence,
        status="analyzed",
        result_json=result_json,
    )

    return ImageAnalyzeResponse(
        status="analyzed",
        ai_probability=result.ai_probability,
        confidence=result.confidence,
        classification=result.classification,
        indicators=result.indicators,
        explanation=result.explanation,
        history_id=history_id,
        provider=result.provider,
    )


# ══════════════════════════════════════════════════════════════════════════════
# NEW: Analysis History + Dashboard (Feature 3)
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/history", response_model=list[HistoryItem])
def get_history_list(filter: str = "all", sort: str = "newest") -> list[HistoryItem]:
    if filter not in ("all", "text", "image"):
        raise HTTPException(status_code=400, detail="filter must be 'all', 'text', or 'image'.")
    if sort not in ("newest", "oldest"):
        raise HTTPException(status_code=400, detail="sort must be 'newest' or 'oldest'.")
    rows = db.list_history(filter_type=filter, sort=sort)  # type: ignore[arg-type]
    return [HistoryItem(**{k: v for k, v in r.items() if k != "result_json"}) for r in rows]


@app.get("/history/{history_id}", response_model=HistoryDetail)
def get_history_detail(history_id: int) -> HistoryDetail:
    row = db.get_history(history_id)
    if row is None:
        raise HTTPException(status_code=404, detail="Analysis not found.")
    return HistoryDetail(**row)


@app.delete("/history/{history_id}")
def delete_history_item(history_id: int) -> dict:
    if not db.delete_history(history_id):
        raise HTTPException(status_code=404, detail="Analysis not found.")
    return {"status": "deleted", "id": history_id}


@app.get("/dashboard", response_model=DashboardStats)
def get_dashboard() -> DashboardStats:
    return DashboardStats(**db.get_dashboard_stats())


# ══════════════════════════════════════════════════════════════════════════════
# NEW: PDF Report Generation (Feature 4)
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/report/{history_id}")
def get_report(history_id: int) -> Response:
    row = db.get_history(history_id)
    if row is None:
        raise HTTPException(status_code=404, detail="Analysis not found.")
    try:
        pdf_bytes = generate_report(row)
    except Exception:
        raise HTTPException(status_code=500, detail="Could not generate the PDF report.")
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="analysis_{history_id}_report.pdf"'},
    )
