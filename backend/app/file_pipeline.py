"""
Page/slide-aware file extraction for the DETECTION pipeline only.

Deliberately independent from the extraction helpers in `app.main`
(`_extract_pdf_text`, `_extract_pptx_text`, `_extract_docx_text`,
`_extract_by_filename`), which remain byte-for-byte unchanged and keep
serving the Summary feature exactly as before. Those helpers flatten a
whole document into one joined string, which is fine for summarization
but throws away page/slide boundaries — this module keeps them, so
multi-page documents (e.g. a PDF with prose on page 1 and Python code on
page 2) can be classified and scored per-unit instead of as one blob.

Does not import from `app.main` (that module imports this one).
"""
from __future__ import annotations

import io
from dataclasses import dataclass

import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation

PAGE_UNIT_EXTENSIONS = {".pdf"}
SLIDE_UNIT_EXTENSIONS = {".pptx", ".ppt"}


@dataclass
class Unit:
    index: int       # true 1-based page/slide number, blanks included
    source: str       # "page 3", "slide 2", "document"
    text: str


def extract_pdf_pages(data: bytes) -> list[str]:
    """One entry per PDF page, in page order, including blank/image-only pages."""
    pages: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            pages.append(text.strip() if text else "")
    return pages


def extract_pptx_slides(data: bytes) -> list[str]:
    """One entry per slide. Bullet lines are kept as separate lines (not
    joined into one sentence) so line-based code detection can see real
    line boundaries."""
    prs = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for slide in prs.slides:
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs).strip()
                    if line:
                        lines.append(line)
        slides.append("\n".join(lines))
    return slides


def _extract_docx_text(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def build_units(filename: str, data: bytes) -> list[Unit]:
    """Dispatch by extension into page/slide/document-granular units."""
    fn = filename.lower()

    if fn.endswith(".pdf"):
        pages = extract_pdf_pages(data)
        return [Unit(index=i + 1, source=f"page {i + 1}", text=t) for i, t in enumerate(pages)]

    if fn.endswith((".pptx", ".ppt")):
        slides = extract_pptx_slides(data)
        return [Unit(index=i + 1, source=f"slide {i + 1}", text=t) for i, t in enumerate(slides)]

    if fn.endswith(".docx"):
        return [Unit(index=1, source="document", text=_extract_docx_text(data))]

    # TXT, or a source-code file — decode as plain text, single unit
    return [Unit(index=1, source="document", text=data.decode("utf-8", errors="ignore"))]
